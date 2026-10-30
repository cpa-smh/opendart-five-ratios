from copy import deepcopy
from decimal import Decimal as D
from unittest.mock import Mock
import json
import pytest
from openpyxl import load_workbook
from sample_data import demo
from pipeline import analyze,load_supplement,save_json,load_snapshot,collect,select_company
from ratios import fraction,calculate
from account_mapper import map_accounts
from dart_client import DartClient,DartError,NoDataError,Company
from excel_report import write_report,SHEETS
from documents import fetch_notes


@pytest.fixture
def inputs(tmp_path):
    s,n=demo(); path=tmp_path/'notes.json'; save_json(n,path)
    return s,load_supplement(path)


def test_five_ratios(inputs):
    obs=analyze(*inputs); r=obs[1]['ratios']
    assert r['operating_margin']['value']==D(170)/1100
    assert r['roe']['value']==D(90)/525
    assert r['ar_turnover']['value']==D(1100)/105
    assert r['debt_ratio']['value']==D(440)/660
    assert r['interest_coverage']['value']==D(170)/22


@pytest.mark.parametrize('n,d',[(None,D(1)),(D(1),None),(D(1),D(0)),(D(1),D(-1))])
def test_invalid_denominators(n,d):
    value,reason=fraction(n,d); assert value is None and reason


def test_negative_profit_and_refund(inputs):
    s,n=inputs
    for row in s['years']['2025']['full']:
        if row['account_id']=='dart_OperatingIncomeLoss': row['thstrm_amount']='-1000000'
    n['00000001',2025]['ecl_expense']=D(-100)
    o=analyze(s,n)[-1]
    assert o['ratios']['operating_margin']['value']<0
    assert o['ratios']['interest_coverage']['value']<0
    assert o['aux']['ecl_to_sales']<0


@pytest.mark.parametrize('year,item',[(2024,'parent_equity'),(2025,'parent_equity'),(2025,'parent_profit')])
def test_roe_fallback_as_pair(inputs,year,item):
    s,n=inputs
    from account_mapper import ACCOUNTS
    s['years'][str(year)]['full']=[r for r in s['years'][str(year)]['full'] if r['account_id']!=ACCOUNTS[item][0][0]]
    r=analyze(s,n)[-1]['ratios']['roe']
    assert r['value']==D(130)/750
    assert '전체' in r['basis']


def test_net_not_deducted_twice_and_rollforward(inputs):
    o=analyze(*inputs)[-1]
    assert o['credit']['net']==D(130000000)
    assert o['credit']['reconciliation']==0
    assert o['rollforward']==0


def test_ecl_not_allowance_change(inputs):
    s,n=inputs; n['00000001',2025]['ecl_expense']=None
    o=analyze(s,n)[-1]
    assert o['aux']['ecl_to_sales'] is None
    assert o['rollforward'] is None


@pytest.mark.parametrize('field,value',[('scope','current_trade'),('fs_div','OFS'),('rcept_no','20260101009999'),('net',D(1)),('allowance',D(999999999))])
def test_invalid_credit_scope_blocks(inputs,field,value):
    s,n=inputs; n['00000001',2025][field]=value
    assert analyze(s,n)[-1]['ratios']['ar_turnover']['value'] is None


def test_rollforward_mismatch(inputs):
    s,n=inputs; n['00000001',2025]['other_movements']=D(5)
    assert analyze(s,n)[-1]['rollforward']==5


def test_other_receivables_not_used(inputs):
    s,n=inputs
    rows=[{'sj_div':'BS','account_id':'ifrs-full_TradeAndOtherCurrentReceivables','account_nm':'매출채권 및 기타채권','thstrm_amount':'5','currency':'KRW','rcept_no':'20260301000001'}]
    assert map_accounts(rows)['receivables'].amount is None


def test_duplicate_candidate_not_summed(inputs):
    rows=inputs[0]['years']['2025']['full']
    extra=deepcopy(rows[0]); extra['thstrm_amount']='1'
    assert map_accounts(rows+[extra])['revenue'].status=='REVIEW'


def test_finance_cost_not_interest():
    rows=[{'sj_div':'IS','account_id':'ifrs-full_FinanceCosts','account_nm':'금융비용','thstrm_amount':'5','currency':'KRW','rcept_no':'20260301000001'}]
    assert map_accounts(rows)['interest'].amount is None


def test_period_mismatch_blocks_ratios(inputs):
    s,n=inputs;s['years']['2025']['period'][1]['thstrm_dt']='2025.01.01 ~ 2025.09.30'
    assert all(r['value'] is None for r in analyze(s,n)[-1]['ratios'].values())


def test_missing_opening_does_not_block_margin(inputs):
    s,n=inputs;s['years']['2024']['period']=[]
    ratios=analyze(s,n)[-1]['ratios']
    assert ratios['roe']['value'] is None
    assert ratios['operating_margin']['value'] is not None


def test_snapshot_identity(inputs,tmp_path):
    s,n=inputs; p=tmp_path/'snapshot.json';save_json(s,p)
    assert load_snapshot(p)['end_year']==2025
    s['years']['2025']['full'][0]['corp_code']='99999999';save_json(s,p)
    with pytest.raises(ValueError):load_snapshot(p)


def test_excel_formulas_charts_and_no_injection(inputs,tmp_path):
    s,n=inputs;s['company']['name']='=HYPERLINK("bad")'
    obs=analyze(s,n);p=write_report(s,obs,tmp_path/'test.xlsx');w=load_workbook(p)
    assert w.sheetnames==SHEETS
    assert len(w['그래프']._charts)==5
    assert w['재무비율계산']['H2'].data_type=='f'
    assert w['재무비율계산']['I2'].value.startswith('=IF')
    assert w['재무비율계산']['D3'].data_type=='f'
    assert w['계정매핑']['B2'].data_type=='s'
    assert len(w['원본재무데이터']['A'])>50
    assert w['그래프']._charts[0].display_blanks=='gap'


@pytest.mark.parametrize('status',['010','011','012','020','013','800','900','901'])
def test_api_status_safe_and_not_retried(status):
    key='a'*40; session=Mock();session.get.return_value=Mock(status_code=200,json=lambda:{'status':status,'message':key})
    c=DartClient(key,session=session,sleep=Mock())
    with pytest.raises(DartError) as e:c.profile('00000001')
    assert key not in str(e.value) and session.get.call_count==1


def test_timeout_retry_safe():
    import requests
    s=Mock();s.get.side_effect=requests.Timeout('secret'*10)
    with pytest.raises(DartError) as e:DartClient('a'*40,session=s,sleep=Mock()).profile('00000001')
    assert s.get.call_count==3 and 'secret' not in str(e.value)


def test_same_name_selection():
    cs=[Company('00000001','동명','000001'),Company('00000002','동명','000002')]
    assert select_company(cs,'동명','000002').code=='00000002'
    assert select_company(cs,'동명',ask=lambda _: '00000001').code=='00000001'


@pytest.mark.parametrize('accept',[True,False])
def test_ofs_consent_and_consistent_scope(inputs,tmp_path,accept):
    s,n=inputs;api=Mock();api.profile.return_value=s['profile']
    def full(code,y,basis):
        if basis=='CFS':raise NoDataError('missing')
        return s['years'][str(y)]['full']
    api.financials.side_effect=full;api.period_accounts.return_value=[]
    result=collect(api,Company('00000001','가상제조','000001'),2025,tmp_path,ask=lambda _:'y' if accept else 'n',download_notes=False)
    assert result['fs_div']==('OFS' if accept else 'CFS')
    assert (any(c.args[2]=='OFS' for c in api.financials.call_args_list))==accept


def test_document_scope_and_zip_safety(tmp_path):
    from io import BytesIO
    from zipfile import ZipFile
    b=BytesIO()
    with ZipFile(b,'w') as z:
        z.writestr('abc_00760.xml','<r>별도</r>');z.writestr('abc_00761.xml','<r>연결</r>')
    api=Mock();api._get.return_value.content=b.getvalue()
    fetch_notes(api,'20260301000001','CFS',tmp_path)
    assert (tmp_path/'20260301000001_CFS.txt').read_text(encoding='utf-8')=='연결'


def test_supplement_missing_source_rejected(tmp_path):
    s,n=demo();n['records'][0]['source_ref']='';p=tmp_path/'n.json';save_json(n,p)
    with pytest.raises(ValueError):load_supplement(p)


def test_non_december_fiscal_period():
    from validation import validate_period,profile_reasons
    s,n=demo();raw=s['years']['2025']
    raw['period'][0]['thstrm_dt']='2025.03.31'
    raw['period'][1]['thstrm_dt']='2024.04.01 ~ 2025.03.31'
    assert validate_period(raw['full'],raw['period'],2025,'CFS',3)['valid']
    assert not profile_reasons({'corp_cls':'Y','acc_mt':'03','induty_code':'261'})


def test_interest_supplement_conflict(inputs):
    s,n=inputs;n['00000001',2025]['interest_expense']=D(1)
    o=analyze(s,n)[-1]
    assert o['ratios']['interest_coverage']['value'] is None
    assert o['ratios']['operating_margin']['value'] is not None


def test_missing_interest_only_does_not_block_others(inputs):
    s,n=inputs;s['years']['2025']['full']=[r for r in s['years']['2025']['full'] if r['account_id']!='ifrs-full_InterestExpense']
    o=analyze(s,n)[-1];assert o['ratios']['interest_coverage']['value'] is None
    assert sum(r['value'] is not None for r in o['ratios'].values())==4


def test_current_net_conflict(inputs):
    s,n=inputs;n['00000001',2025]['current_net']='1'
    o=analyze(s,n)[-1]
    assert o['credit']['net'] is None and o['credit']['reconciliation'] is None


def test_png_and_missing_values(inputs,tmp_path):
    from charts import generate
    obs=analyze(*inputs)
    for o in obs[1:]:o['ratios']['ar_turnover']['value']=None
    try:
        paths=generate(obs[1:],tmp_path,True)
    except ValueError:
        pytest.skip('한글 글꼴 없는 실행 환경')
    assert len(paths)==5 and all(p.stat().st_size>1000 for p in paths)
