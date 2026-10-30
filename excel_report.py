"""openpyxl 보고서. Excel 수식과 Python 결과를 나란히 둔다."""
from decimal import Decimal
from pathlib import Path
import json
from openpyxl import Workbook
from openpyxl.chart import LineChart, Reference
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.formatting.rule import CellIsRule
from openpyxl.workbook.properties import CalcProperties
from ratios import METRICS

SHEETS = ['분석요약','원본재무데이터','계정매핑','재무비율계산','매출채권_ECL','검증결과','그래프']


def safe(value):
    if isinstance(value,Decimal):
        return float(value)
    if isinstance(value,(dict,list,tuple)):
        value=json.dumps(value,ensure_ascii=False,default=str)
    if isinstance(value,str) and value.startswith(('=','+','-','@')):
        return "'"+value
    return value


def append(sheet, row):
    sheet.append([safe(v) for v in row])


def write_report(snapshot, observations, path):
    wb=Workbook(); wb.remove(wb.active)
    for name in SHEETS: wb.create_sheet(name)
    wb.calculation=CalcProperties(calcId=191029, fullCalcOnLoad=True, forceFullCalc=True)
    title=('[가상 데모] ' if snapshot['synthetic'] else '')+snapshot['company']['name']
    summary=wb['분석요약']; rawws=wb['원본재무데이터']; mapping=wb['계정매핑']; calc=wb['재무비율계산']; cr=wb['매출채권_ECL']; checks=wb['검증결과']; graph=wb['그래프']
    append(summary,[title+' 재무비율 분석'])
    append(summary,[f"{snapshot['end_year']-2}~{snapshot['end_year']} | {snapshot['fs_div']} | 금액: 원, 비율: %, 배"])
    append(summary,['회전율은 전체 매출액을 외상매출액 대신 사용합니다. 업종별 정상·위험 판정은 하지 않습니다.'])
    append(summary,['지표','단위',*range(snapshot['end_year']-2,snapshot['end_year']+1),'해석'])
    rawkeys=['account_nm','thstrm_amount','currency','sj_div','account_id','rcept_no']
    otherkeys=sorted({k for o in observations for group in ('full','period') for r in o['raw'][group] for k in r}-set(rawkeys))
    rawkeys+=otherkeys
    append(rawws,['조회연도','구분','기업','재무제표',*rawkeys,'공시 URL'])
    append(mapping,['연도','기업','재무제표','항목','금액(원)','계정 ID','계정명','상태','선택 근거','접수번호','통화','공시 URL','후보'])
    append(calc,['연도','지표','분자(원)','기초 분모(원)','기말 분모(원)','사용 분모(원)','Python 결과','Excel 결과','차이','계산 불가 사유','기준','계산 허용'])
    append(cr,['연도','총 매출채권','손실충당금','순액','대손·ECL 비용','제각','기타 변동','설정률','ECL/매출','출처 구분','주석 위치','공시 URL','상태/사유'])
    append(checks,['연도','검증','차이(원)','결과/사유'])
    mapcells={}; creditrows={}; calccells={}
    for o in observations:
        year=o['year']; note=o['note'] or {}; credit=o['credit']
        for group in ('full','period'):
            for raw in o['raw'][group]:
                append(rawws,[year,group,snapshot['company']['name'],snapshot['fs_div'],*[raw.get(k) for k in rawkeys],'https://dart.fss.or.kr/dsaf001/main.do?rcpNo='+raw.get('rcept_no','')])
        for key,m in o['mapping'].items():
            append(mapping,[year,snapshot['company']['name'],o['fs_div'],key,m.amount,m.account_id,m.account_nm,m.status,m.method+': '+m.reason,m.receipt,m.currency,'https://dart.fss.or.kr/dsaf001/main.do?rcpNo='+m.receipt,m.candidates])
            mapcells[year,key]=f"'계정매핑'!E{mapping.max_row}"
        if o['interest_source']!='API':
            append(mapping,[year,snapshot['company']['name'],o['fs_div'],'interest_effective',o['values']['interest'],'보완 입력','손익 이자비용',o['interest_source'],note.get('source_ref',''),note.get('rcept_no',''),'KRW',note.get('source_url','')])
            mapcells[year,'interest']=f"'계정매핑'!E{mapping.max_row}"
        append(cr,[year,credit['gross'],credit['allowance'],credit['net'],credit['ecl_expense'],credit['writeoffs'],credit['other_movements'],None,None,note.get('origin','미확인'),credit['source_ref'],credit['source_url'],credit['status']+': '+credit['reason']])
        row=cr.max_row; creditrows[year]=row
        mapcells[year,'net']=f"'매출채권_ECL'!D{row}"
        cr.cell(row,8,f'=IF(AND(ISNUMBER(B{row}),B{row}>0,ISNUMBER(C{row})),C{row}/B{row},"")' if o['eligible'] else '=""')
        rev=mapcells[year,'revenue']
        cr.cell(row,9,f'=IF(AND(ISNUMBER(E{row}),ISNUMBER({rev}),{rev}>0),E{row}/{rev},"")' if o['eligible'] else '=""')
        for check in o['checks']:
            append(checks,[year,check['check'],check['difference'],check['status']])
        append(checks,[year,'기간·통화 확인',None,'PASS' if o['period']['valid'] else '; '.join(o['period']['reasons'])])
        for issue in o['issues']: append(checks,[year,'수집/대사',None,issue])
        append(checks,[year,'매출채권 총액-충당금-순액',credit['reconciliation'],'SKIP' if credit['reconciliation'] is None else 'PASS' if credit['reconciliation']==0 else 'FAIL'])
        if 'rollforward' in o:
            append(checks,[year,'충당금 롤포워드',o['rollforward'],'SKIP' if o['rollforward'] is None else 'PASS' if o['rollforward']==0 else 'FAIL'])
        for key,m in o['mapping'].items():
            if m.status!='OK': append(checks,[year,'계정 '+key,None,m.status+': '+m.reason])
    for o in observations[1:]:
        year=o['year']
        for key,(label,unit,_) in METRICS.items():
            r=o['ratios'][key]; parent=r['basis']=='지배기업 소유주'
            nkey={'operating_margin':'operating_profit','roe':'parent_profit' if parent else 'profit','ar_turnover':'revenue','debt_ratio':'liabilities','interest_coverage':'operating_profit'}[key]
            dkey={'operating_margin':'revenue','roe':'parent_equity' if parent else 'equity','ar_turnover':'net','debt_ratio':'equity','interest_coverage':'interest'}[key]
            averaged=key in ('roe','ar_turnover')
            append(calc,[year,label,None,None,None,None,r['value'],None,None,r['reason'],r['basis'],r['value'] is not None])
            row=calc.max_row; calccells[year,key]=row
            def link(col,source):
                calc.cell(row,col,f'=IF(ISNUMBER({source}),{source},"")')
            link(3,mapcells[year,nkey]); link(5,mapcells[year,dkey])
            if averaged: link(4,mapcells[year-1,dkey])
            calc.cell(row,6,f'=IF(AND(ISNUMBER(D{row}),ISNUMBER(E{row})),(D{row}+E{row})/2,"")' if averaged else f'=E{row}')
            calc.cell(row,8,f'=IF(AND(L{row},ISNUMBER(C{row}),ISNUMBER(F{row}),F{row}>0),C{row}/F{row},"")')
            calc.cell(row,9,f'=IF(AND(ISNUMBER(G{row}),ISNUMBER(H{row})),H{row}-G{row},"")')
            for c in (7,8): calc.cell(row,c).number_format='0.00%' if unit=='%' else '0.00"배"'
    meanings=['매출 대비 영업이익','평균 자본 대비 순이익; 지배주주/전체 범위는 계산 시트 참조','매출대금 회전 속도의 대용 지표','자본 대비 부채 규모','영업이익으로 손익 이자비용을 감당하는 정도']
    for i,(key,(label,unit,_)) in enumerate(METRICS.items(),5):
        first,last=observations[1]['ratios'][key]['value'],observations[-1]['ratios'][key]['value']
        if first is not None and last is not None:
            delta=(last-first)*(100 if unit=='%' else 1)
            trend=f"첫해 대비 {delta:+.2f}{'%p' if unit=='%' else '배'}. "
        else:
            trend='자료 미확인으로 첫해·마지막해 비교 불가. '
        append(summary,[label,unit,None,None,None,trend+meanings[i-5]])
        for col,o in enumerate(observations[1:],3):
            row=calccells[o['year'],key]
            summary.cell(i,col,f'=IF(ISNUMBER(\'재무비율계산\'!G{row}),\'재무비율계산\'!G{row},"미확인")')
            summary.cell(i,col).number_format='0.00%' if unit=='%' else '0.00"배"'
    append(summary,[])
    append(summary,['Python 결과를 표시합니다. Excel 계산식·차이는 재무비율계산 시트에서 확인하세요.'])
    append(summary,['Excel 수식은 열 때 재계산됩니다. 입력 변경 시 계산 허용 조건도 확인하고 Python을 다시 실행하세요.'])
    append(summary,['ECL은 공시된 대손상각(환입)을 사용하며 충당금 잔액 증감으로 추정하지 않습니다.'])
    append(summary,['각 연도 조회 시점의 당기 금액 기준. 후속 보고서의 비교표시 금액으로 소급 대체하지 않습니다.'])
    append(graph,[title+' 추세 그래프'])
    append(graph,['연도',*[x[0] for x in METRICS.values()]])
    for idx,o in enumerate(observations[1:],3):
        graph.cell(idx,1,str(o['year']))
        for col,key in enumerate(METRICS,2):
            row=calccells[o['year'],key]
            graph.cell(idx,col,f'=IF(ISNUMBER(\'재무비율계산\'!H{row}),\'재무비율계산\'!H{row},"")')
            graph.cell(idx,col).number_format='0.00%' if METRICS[key][1]=='%' else '0.00'
    for i,(key,(label,unit,_)) in enumerate(METRICS.items()):
        chart=LineChart(); chart.title=label; chart.y_axis.title=unit; chart.x_axis.title='사업연도'
        chart.add_data(Reference(graph,min_col=i+2,min_row=2,max_row=5),titles_from_data=True)
        chart.set_categories(Reference(graph,min_col=1,min_row=3,max_row=5))
        chart.display_blanks='gap'; chart.height=8; chart.width=17; chart.legend=None
        chart.series[0].smooth=False
        chart.y_axis.numFmt='0.0%' if unit=='%' else '0.0'
        chart.series[0].graphicalProperties.line.solidFill='205493'
        graph.add_chart(chart,f'{"A" if i%2==0 else "J"}{8+(i//2)*17}')
    for ws in wb:
        ws.sheet_view.showGridLines=False
        ws.freeze_panes='B2' if ws.title!='분석요약' else None
        for row in ws:
            for cell in row:
                cell.font=Font(name='맑은 고딕',size=10,color='1F2937')
                cell.alignment=Alignment(vertical='center')
                if isinstance(cell.value,(float,int)) and not isinstance(cell.value,bool) and cell.number_format=='General':
                    cell.number_format='#,##0;(#,##0);"-"'
        headerrow=4 if ws==summary else 2 if ws==graph else 1
        for cell in ws[headerrow]:
            cell.fill=PatternFill('solid',fgColor='173F5F'); cell.font=Font(name='맑은 고딕',size=10,bold=True,color='FFFFFF')
            cell.alignment=Alignment(horizontal='center',vertical='center',wrap_text=True)
        ws.row_dimensions[headerrow].height=32
        for col in range(1,ws.max_column+1):
            from openpyxl.utils import get_column_letter
            ws.column_dimensions[get_column_letter(col)].width=22
        ws.column_dimensions['A'].width=15
        if ws not in (summary,graph): ws.auto_filter.ref=ws.dimensions
        for i in range(2,ws.max_row+1): ws.row_dimensions[i].height=22
        for cell in ws['A'][1:]:
            if isinstance(cell.value,int): cell.number_format='0'
    for cell in summary[4][2:5]: cell.number_format='0'
    summary.column_dimensions['B'].width=8
    for col in ('C','D','E'): summary.column_dimensions[col].width=18
    for col in range(1,19): graph.column_dimensions[get_column_letter(col)].width=10
    summary.column_dimensions['A'].width=24; summary.column_dimensions['F'].width=74
    for row in range(5,10):
        summary.cell(row,6).alignment=Alignment(wrap_text=True,vertical='center')
        summary.row_dimensions[row].height=34
    for i in (1,2,3,11,12,13,14): summary.row_dimensions[i].height=25
    rawws.column_dimensions['E'].width=42; rawws.column_dimensions['F'].width=25
    for index,name in enumerate(rawkeys,5):
        rawws.column_dimensions[get_column_letter(index)].width=42 if name in ('account_nm','account_id','account_detail') else 25
    rawws.column_dimensions[get_column_letter(rawws.max_column)].width=65
    for col in ('F','G','I','L','M'): mapping.column_dimensions[col].width=45
    calc.column_dimensions['J'].width=68; calc.column_dimensions['K'].width=36
    calc['I1']='Excel - Python'
    for row in calc.iter_rows(min_row=2):
        for cell in row[2:6]: cell.number_format='#,##0;(#,##0);"-"'
        row[8].number_format='0.0000000000'
    cr.column_dimensions['K'].width=80; cr.column_dimensions['L'].width=65; cr.column_dimensions['M'].width=65
    for row in cr.iter_rows(min_row=2):
        row[7].number_format=row[8].number_format='0.000%'
    checks.column_dimensions['B'].width=38; checks.column_dimensions['D'].width=95
    checks.conditional_formatting.add(f'C2:C{checks.max_row}',CellIsRule(operator='notEqual',formula=['0'],fill=PatternFill('solid',fgColor='FEE2E2')))
    Path(path).parent.mkdir(parents=True,exist_ok=True); wb.save(path)
    return Path(path)
